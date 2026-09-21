"""Export a finished run to an .xlsx workbook for the operations team.

Reads the same data the dashboard does, so the sheet can never disagree with it.
Makes no API calls: every value here was already decided by the run.

    python scripts/export_excel.py                  # -> results/ladinglens.xlsx
    python scripts/export_excel.py --sample         # the ?limit= run instead
    python scripts/export_excel.py -o handover.xlsx

Two sheets:
  Inbox      one row per email -- what it was classified as and what happened
  Mismatches one row per flagged field, with the SI and BL values side by side
"""
import argparse
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from dotenv import load_dotenv  # noqa: E402

load_dotenv(ROOT / ".env")

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from app import api  # noqa: E402
from app.paths import results_file  # noqa: E402

FIELD_LABEL = {
    "shipper": "Shipper",
    "consignee": "Consignee",
    "notify_party": "Notify party",
    "port_of_loading": "Port of loading",
    "port_of_discharge": "Port of discharge",
    "container_count": "Container count",
    "gross_weight_kg": "Gross weight (kg)",
}
CATEGORY_LABEL = {
    "BL_COMPARISON": "Comparison request",
    "SI_REQUEST": "New SI request",
    "INVOICE_QUERY": "Invoice query",
    "GENERAL": "General",
    "SPAM": "Spam",
}
REASON_LABEL = {
    "unreadable": "Document unreadable",
    "missing_attachment": "Attachment missing",
    "missing_value": "Required value missing",
    "wrong_doc_type": "Wrong document type",
    "low_confidence": "Low classification confidence",
    "processing_error": "Processing failed (API, not the document)",
}

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
# Status colours, not decoration: a reviewer scans for the red rows.
STATUS_FILL = {
    "MISMATCH": PatternFill("solid", fgColor="FDE2E1"),
    "NEEDS_REVIEW": PatternFill("solid", fgColor="FEF3C7"),
}


def _style_header(ws, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _clean(value) -> str:
    """Flatten to one line -- addresses arrive with newlines that wreck row heights."""
    if value is None:
        return ""
    return " ".join(str(value).split())


def build(sample: bool) -> Workbook:
    report = api.report(scope="sample" if sample else "auto")
    if not report["emails"]:
        sys.exit("No results found. Run the pipeline first (POST /run).")

    wb = Workbook()

    inbox = wb.active
    inbox.title = "Inbox"
    inbox.append([
        "Email", "From", "Subject", "Classified as", "Confidence",
        "Result", "Why escalated", "Fields flagged",
    ])
    mismatches = wb.create_sheet("Mismatches")
    mismatches.append([
        "Email", "From", "Subject", "Field",
        "Shipping Instruction", "Draft Bill of Lading",
    ])

    for row in report["emails"]:
        status = row["status"]
        flagged = ", ".join(FIELD_LABEL.get(f, f) for f in row["defect_fields"])
        inbox.append([
            row["email_id"],
            row.get("from", ""),
            _clean(row.get("subject")),
            CATEGORY_LABEL.get(row["category"], row["category"]),
            row.get("confidence"),
            {"OK": "No mismatch", "MISMATCH": "Mismatch", "NEEDS_REVIEW": "Needs review"}[status],
            REASON_LABEL.get(row.get("review_reason"), ""),
            flagged,
        ])
        if fill := STATUS_FILL.get(status):
            for cell in inbox[inbox.max_row]:
                cell.fill = fill

        if not row["has_defect"]:
            continue
        # Only flagged emails need the per-field detail, so only those are fetched.
        case = api.email_case(row["email_id"], scope=report["scope"])
        for field, pair in (case["record"].get("mismatches") or {}).items():
            mismatches.append([
                row["email_id"],
                row.get("from", ""),
                _clean(row.get("subject")),
                FIELD_LABEL.get(field, field),
                _clean(pair.get("si")),
                _clean(pair.get("bl")),
            ])

    _style_header(inbox, [12, 28, 52, 20, 11, 14, 34, 30])
    _style_header(mismatches, [12, 28, 42, 20, 46, 46])
    return wb


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--sample", action="store_true", help="export the ?limit= run")
    ap.add_argument("-o", "--out", help="output path (default: results/ladinglens.xlsx)")
    args = ap.parse_args()

    wb = build(args.sample)
    out = Path(args.out) if args.out else results_file(
        "ladinglens.sample.xlsx" if args.sample else "ladinglens.xlsx"
    )
    wb.save(out)
    print(f"wrote {out}")
    for name in wb.sheetnames:
        print(f"  {name}: {wb[name].max_row - 1} rows")


if __name__ == "__main__":
    main()
